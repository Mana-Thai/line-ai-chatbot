# -*- coding: utf-8 -*-
"""agent.py の自己診断。APIキーなしで、ツールがこのPCで動くかを確認する。

    python selftest.py

一時フォルダにサンプル(Excel/Word/PDF/画像)を作って、agent.py のツール関数を
直接呼ぶ。API通信は行わないので課金されない。導入直後や、環境を変えたときに実行する。
FAIL が1件でもあれば exit 1(未導入のライブラリは SKIP 扱い)。
"""
from __future__ import annotations

import shutil
import sys
import tempfile
import types
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import agent as A


class Report:
    def __init__(self) -> None:
        self.ok = self.fail = self.skip = 0

    def check(self, label: str, fn, expect: str | None = None) -> None:
        try:
            got = fn()
        except Exception as e:
            got = f"<{type(e).__name__}: {e}>"
        passed = (expect in str(got)) if expect is not None else True
        self.ok, self.fail = self.ok + passed, self.fail + (not passed)
        mark = "\033[32mOK  \033[0m" if passed else "\033[31mFAIL\033[0m"
        detail = str(got).replace("\n", " | ")[:100]
        print(f"  [{mark}] {label}" + ("" if passed else f"  -> {detail}"))

    def skipped(self, label: str, why: str) -> None:
        self.skip += 1
        print(f"  [\033[33mSKIP\033[0m] {label}  ({why})")


def has(module: str) -> bool:
    try:
        __import__(module)
        return True
    except ImportError:
        return False


def make_pdf(path: Path, text: str) -> None:
    """xref付きの最小PDFを組み立てる(外部ライブラリ不要)。"""
    stream = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode()
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>stream\n" % len(stream) + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj" % i + body + b"endobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer<</Root 1 0 R/Size %d>>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.write_bytes(bytes(out))


def main() -> int:
    root = Path(tempfile.mkdtemp(prefix="agent-selftest-"))
    (root / "sub").mkdir()
    g = A.Guard([root], auto_approve=True, read_only=False)
    r = Report()

    print(f"\n作業フォルダ: {root}\n")
    print("■ フォルダ制限")
    r.check("相対パスは許可フォルダ基準で解決", lambda: str(g.resolve("sub")), str(root / "sub"))
    r.check("外部の絶対パスを拒否", lambda: g.resolve(str(Path.home())), "許可フォルダの外")
    r.check("..での脱出を拒否", lambda: g.resolve("sub/../../.."), "許可フォルダの外")

    print("■ ファイルの閲覧・編集")
    memo = root / "memo.txt"
    r.check("作成", lambda: A.tool_text_editor(
        {"command": "create", "path": str(memo), "file_text": "一行目\n二行目\n三行目\n"}, g), "書き込みました")
    r.check("閲覧(行番号つき)", lambda: A.tool_text_editor({"command": "view", "path": str(memo)}, g), "2\t二行目")
    r.check("範囲指定の閲覧", lambda: A.tool_text_editor(
        {"command": "view", "path": str(memo), "view_range": [3, 3]}, g), "3\t三行目")
    r.check("フォルダの一覧", lambda: A.tool_text_editor({"command": "view", "path": str(root)}, g), "[DIR ] sub/")
    r.check("部分置換", lambda: A.tool_text_editor(
        {"command": "str_replace", "path": str(memo), "old_str": "二行目", "new_str": "2行目"}, g), "置換しました")
    r.check("一致しない置換は拒否", lambda: A.tool_text_editor(
        {"command": "str_replace", "path": str(memo), "old_str": "無い", "new_str": "x"}, g), "見つかりません")
    r.check("行の挿入", lambda: A.tool_text_editor(
        {"command": "insert", "path": str(memo), "insert_line": 1, "insert_text": "挿入行"}, g), "挿入しました")
    r.check("上書き時に .bak を残す", lambda: (A.tool_text_editor(
        {"command": "create", "path": str(memo), "file_text": "上書き\n"}, g),
        (root / "memo.txt.bak").exists())[1], "True")

    print("■ 検索")
    (root / "sub" / "data.csv").write_text("名前,金額\n田中,1200\n", encoding="utf-8")
    r.check("ファイル名で探す", lambda: A.tool_search({"glob": "*.csv"}, g), "data.csv")
    r.check("中身で探す", lambda: A.tool_search({"pattern": "田中"}, g), "data.csv:2")
    r.check("該当なしを報告", lambda: A.tool_search({"pattern": "存在しない語句zzz"}, g), "該当なし")

    print("■ 文書ファイルの読み取り")
    if has("openpyxl"):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.title = "見積"
        wb.active.append(["品名", "数量"])
        wb.active.append(["Tシャツ", 20])
        wb.save(root / "book.xlsx")
        r.check("Excel", lambda: A.tool_read_document({"path": str(root / "book.xlsx")}, g), "Tシャツ\t20")
    else:
        r.skipped("Excel", "pip install openpyxl")
    if has("docx"):
        import docx
        d = docx.Document()
        d.add_paragraph("請求書 合計 17000円")
        d.save(root / "doc.docx")
        r.check("Word", lambda: A.tool_read_document({"path": str(root / "doc.docx")}, g), "17000円")
    else:
        r.skipped("Word", "pip install python-docx")
    if has("pypdf"):
        make_pdf(root / "invoice.pdf", "Invoice No 12345")
        r.check("PDF", lambda: A.tool_read_document({"path": str(root / "invoice.pdf")}, g), "12345")
        r.check("PDFのページ指定", lambda: A.parse_pages("2-4", 10), "[1, 2, 3]")
    else:
        r.skipped("PDF", "pip install pypdf")
    if has("PIL"):
        from PIL import Image
        Image.new("RGB", (3000, 2000), (200, 30, 30)).save(root / "photo.png")
        r.check("画像を縮小して渡す", lambda: A.tool_read_document({"path": str(root / "photo.png")}, g)[0]["text"],
                "2000x1333")
    else:
        r.skipped("画像", "pip install pillow")
    (root / "old.xls").write_bytes(b"dummy")
    r.check("旧形式は理由を返す", lambda: A.tool_read_document({"path": str(root / "old.xls")}, g), "旧形式")

    print("■ シェル")
    r.check("実行と出力", lambda: A.tool_shell({"command": "echo こんにちは"}, g), "こんにちは")
    r.check("終了コードを返す", lambda: A.tool_shell({"command": "exit 3"}, g), "exit code: 3")
    r.check("許可外フォルダでの実行を拒否", lambda: A.tool_shell(
        {"command": "echo x", "cwd": str(Path.home())}, g), "許可フォルダの外")

    print("■ 安全装置")
    ro = A.Guard([root], auto_approve=True, read_only=True)
    r.check("読み取り専用は書き込みを拒否", lambda: A.tool_text_editor(
        {"command": "create", "path": str(root / "x.txt"), "file_text": "y"}, ro), "読み取り専用")
    r.check("読み取り専用でも閲覧は可能", lambda: A.tool_text_editor({"command": "view", "path": str(memo)}, ro), "上書き")

    msgs = [{"role": "user", "content": [{"type": "tool_result", "content": [
        {"type": "image", "source": {"type": "base64", "data": f"img{i}"}}]}]} for i in range(5)]
    A.prune_screenshots(msgs, keep=2)
    r.check("古い画像を履歴から落とす",
            lambda: [m["content"][0]["content"][0]["type"] for m in msgs],
            "['text', 'text', 'text', 'image', 'image']")

    log_path = root / "log.jsonl"
    args = types.SimpleNamespace(model="m", effort="high", max_steps=1, no_compact=False)
    ag = A.Agent(None, args, g, None, "sys", log_path)
    ag.log("shell", {"command": "dir", "他": "記録しない"}, "ok")
    import json
    rec = json.loads(log_path.read_text(encoding="utf-8").splitlines()[0])
    r.check("操作ログを残す", lambda: rec["tool"] + "/" + rec["input"]["command"], "shell/dir")
    r.check("対象外の入力は記録しない", lambda: list(rec["input"]), "['command']")

    print("■ 長い会話の自動要約(400なら自動でオフ)")
    calls: list[dict] = []

    class _Stream:
        def __init__(self, kw): self.kw = kw
        def __enter__(self):
            calls.append(self.kw)
            if "context_management" in self.kw:
                import anthropic, httpx
                raise anthropic.BadRequestError(
                    "context_management: compact_20260112 is not supported",
                    response=httpx.Response(400, request=httpx.Request("POST", "https://x")), body=None)
            return self
        def __exit__(self, *a): return False
        def __iter__(self): return iter([])
        def get_final_message(self): return types.SimpleNamespace(stop_reason="end_turn", content=[])

    fake = types.SimpleNamespace(beta=types.SimpleNamespace(
        messages=types.SimpleNamespace(stream=lambda **kw: _Stream(kw))))
    ag2 = A.Agent(fake, args, g, None, "sys", None)
    got = ag2.call_model([{"role": "user", "content": "hi"}])
    r.check("非対応なら要約を外して再試行", lambda: len(calls) == 2 and "context_management" not in calls[1], "True")
    r.check("再試行後も応答を返す", lambda: got is not None, "True")

    print("■ ツール構成")
    r.check("既定は4ツール", lambda: [t.get("name") for t in A.build_tools(None)],
            "['str_replace_based_edit_tool', 'search_files', 'read_document', 'shell']")

    shutil.rmtree(root, ignore_errors=True)
    print(f"\n結果: OK {r.ok} / FAIL {r.fail} / SKIP {r.skip}")
    if r.fail:
        print("\033[31mFAIL があります。README のセットアップを確認してください。\033[0m")
    return 1 if r.fail else 0


if __name__ == "__main__":
    sys.exit(main())
